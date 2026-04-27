from __future__ import annotations
import io
from zipfile import BadZipFile
import openpyxl
import pandas as pd
from config.settings import INPUT_FILE, WORKING_FILE


def current_workbook_path():
    if WORKING_FILE.exists():
        try:
            openpyxl.load_workbook(WORKING_FILE, read_only=True)
            return WORKING_FILE
        except Exception:
            try:
                WORKING_FILE.unlink()
            except Exception:
                pass
    return INPUT_FILE


def load_raw_sheets() -> dict[str, pd.DataFrame]:
    path = current_workbook_path()
    xls = pd.ExcelFile(path)
    sheets = {}
    for name in xls.sheet_names:
        if name == "Packing&TQM":
            sheets[name] = pd.read_excel(path, sheet_name=name, header=None)
        else:
            sheets[name] = pd.read_excel(path, sheet_name=name)
    return sheets


def _write_df(ws, df: pd.DataFrame):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    for c_idx, col in enumerate(df.columns, start=1):
        ws.cell(1, c_idx).value = col
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value


def build_sheet_frames(state: dict) -> dict[str, pd.DataFrame]:
    return {
        "Packing": state["ppc_df"].drop(columns=[c for c in ["row_order"] if c in state["ppc_df"].columns]),
        "Capacity": state["capacity_df"].drop(columns=[c for c in ["row_order"] if c in state["capacity_df"].columns]),
        "Coating": state["coating_df"].drop(columns=[c for c in ["row_order"] if c in state["coating_df"].columns]),
        "CS": state["cs_df"].drop(columns=[c for c in ["row_order"] if c in state["cs_df"].columns]),
        "Hydro ": state["hydro_df"].drop(columns=[c for c in ["row_order"] if c in state["hydro_df"].columns]),
        "Rugs": state["master_df"].drop(columns=[c for c in ["row_order", "row_id"] if c in state["master_df"].columns]),
    }


def _clear_sheet(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None


def _write_packintqm_sheet(ws, state: dict):
    _clear_sheet(ws)
    helper_df = state.get("packing_tqm_helper_df", pd.DataFrame()).copy()
    draft_df = state["master_df"][state["master_df"]["Section"].isin(["Packing  Line MP", "Packing Support MP", "TQM", "TQM Other Support"])].copy()
    helper_df = helper_df.drop(columns=[c for c in ["row_order", "row_id"] if c in helper_df.columns], errors='ignore')
    for c_idx, col in enumerate(helper_df.columns, start=1):
        ws.cell(2, c_idx).value = col
    for r_idx, row in enumerate(helper_df.itertuples(index=False), start=3):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value
    start = len(helper_df.index) + 10
    draft_df = draft_df.drop(columns=[c for c in ["row_order", "row_id"] if c in draft_df.columns], errors='ignore')
    for c_idx, col in enumerate(draft_df.columns, start=1):
        ws.cell(start, c_idx).value = col
    for r_idx, row in enumerate(draft_df.itertuples(index=False), start=start+1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value


def _load_base_workbook():
    path = current_workbook_path()
    try:
        return openpyxl.load_workbook(path)
    except (BadZipFile, FileNotFoundError, OSError):
        return openpyxl.load_workbook(INPUT_FILE)


def save_snapshot(state: dict):
    WORKING_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = _load_base_workbook()
    for sheet_name, df in build_sheet_frames(state).items():
        if sheet_name not in wb.sheetnames:
            continue
        _write_df(wb[sheet_name], df)
    if "Packing&TQM" in wb.sheetnames:
        _write_packintqm_sheet(wb["Packing&TQM"], state)
    wb.save(WORKING_FILE)
    return WORKING_FILE


def workbook_bytes(state: dict) -> bytes:
    wb = _load_base_workbook()
    for sheet_name, df in build_sheet_frames(state).items():
        if sheet_name not in wb.sheetnames:
            continue
        _write_df(wb[sheet_name], df)
    if "Packing&TQM" in wb.sheetnames:
        _write_packintqm_sheet(wb["Packing&TQM"], state)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def reset_working_file():
    if WORKING_FILE.exists():
        WORKING_FILE.unlink()
