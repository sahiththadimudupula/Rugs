from __future__ import annotations
import io
import openpyxl
import pandas as pd
from config.settings import INPUT_FILE, WORKING_FILE

def current_workbook_path():
    return WORKING_FILE if WORKING_FILE.exists() else INPUT_FILE

def load_raw_sheets() -> dict[str, pd.DataFrame]:
    xls = pd.ExcelFile(current_workbook_path())
    return {name: pd.read_excel(current_workbook_path(), sheet_name=name) for name in xls.sheet_names}

def _write_df(ws, df: pd.DataFrame):
    for c_idx, col in enumerate(df.columns, start=1):
        ws.cell(1, c_idx).value = col
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value

def build_sheet_frames(state: dict) -> dict[str, pd.DataFrame]:
    return {
        "Packing": state["ppc_df"].drop(columns=["row_order"]),
        "Capacity": state["capacity_df"].drop(columns=["row_order"]),
        "Coating": state["coating_df"].drop(columns=["row_order"]),
        "CS": state["cs_df"].drop(columns=["row_order"]),
        "Hydro ": state["hydro_df"].drop(columns=["row_order"]),
        "Rugs": state["master_df"].drop(columns=["row_order","row_id"]),
    }

def save_snapshot(state: dict):
    wb = openpyxl.load_workbook(current_workbook_path())
    for sheet_name, df in build_sheet_frames(state).items():
        if sheet_name not in wb.sheetnames:
            continue
        _write_df(wb[sheet_name], df)
    wb.save(WORKING_FILE)
    return WORKING_FILE

def workbook_bytes(state: dict) -> bytes:
    wb = openpyxl.load_workbook(current_workbook_path())
    for sheet_name, df in build_sheet_frames(state).items():
        if sheet_name not in wb.sheetnames:
            continue
        _write_df(wb[sheet_name], df)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

def reset_working_file():
    if WORKING_FILE.exists():
        WORKING_FILE.unlink()
